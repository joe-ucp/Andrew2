"""Post-review, bounded HGNC/Ensembl identifier sensitivity; never changes primary.

All source symbols are retained. Exact approved symbols take precedence, otherwise
only a unique previous/alias-symbol target is accepted. Expression joins use the
source Ensembl GeneID, with explicit symbol disagreement and collision exclusions.
No main-analysis construction code is imported. Inference uses the shared tested
exact conditional helper; mapping and universe construction are independent.
"""
from collections import Counter, defaultdict
import csv
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import openpyxl
import numpy as np
import pandas as pd
from exact_stats import exact_conditional

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
HGNC = ROOT / "data/source_documents/hgnc_complete_set.txt"
HGNC_HASH = "12363916be36f8f6f91879b0faade4eeaac973722de7136ec61329b77f2c2f91"

def number(v):
    try:
        return math.isfinite(float(v))
    except (ValueError, TypeError):
        return False

def write_csv(path, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)

def main():
    assert hashlib.sha256(HGNC.read_bytes()).hexdigest() == HGNC_HASH
    manifest = json.loads((ROOT / "evidence/input_manifest.json").read_text())
    for entry in manifest["inputs"]:
        if entry["id"] in {"D1", "D2"}:
            assert hashlib.sha256((ROOT / entry["path"]).read_bytes()).hexdigest() == entry["sha256"], entry["id"]
    with HGNC.open(encoding="utf-8", newline="") as f:
        hgnc_rows = [r for r in csv.DictReader(f, delimiter="\t") if r["status"] == "Approved"]
    by_id = {r["hgnc_id"]: r for r in hgnc_rows}
    current, historic, ensembl = defaultdict(set), defaultdict(set), defaultdict(set)
    for row in hgnc_rows:
        h = row["hgnc_id"]
        current[row["symbol"]].add(h)
        for field in ["prev_symbol", "alias_symbol"]:
            for s in row[field].split("|"):
                if s: historic[s].add(h)
        for e in row["ensembl_gene_id"].split("|"):
            if e: ensembl[e].add(h)
    def resolve(s):
        targets = current.get(s, set())
        if targets:
            return targets, "approved_symbol" if len(targets) == 1 else "ambiguous_approved_symbol"
        targets = historic.get(s, set())
        return targets, ("unique_previous_or_alias_symbol" if len(targets) == 1 else "ambiguous_previous_or_alias_symbol" if targets else "unmapped_symbol")
    ledger = []
    def symbol_map(source, symbols):
        batch = []
        for symbol in sorted(symbols):
            targets, method = resolve(symbol)
            hid = next(iter(targets)) if len(targets) == 1 else ""
            batch.append({"source": source, "source_symbol": symbol, "source_GeneID": "", "source_row": "", "mapping_method": method,
                "candidate_HGNC_ids": ";".join(sorted(targets)), "HGNC_id": hid, "approved_symbol": by_id[hid]["symbol"] if hid else "", "included_mapping": bool(hid), "exclusion_reason": "" if hid else method})
        collisions = Counter(r["HGNC_id"] for r in batch if r["included_mapping"])
        for r in batch:
            if r["included_mapping"] and collisions[r["HGNC_id"]] > 1:
                r["included_mapping"] = False; r["exclusion_reason"] = "multiple_source_symbols_to_one_HGNC_id"
        ledger.extend(batch)
        return {r["source_symbol"]: r["HGNC_id"] for r in batch if r["included_mapping"]}
    with (RAW / "Supplementary_Data_09.CRISPR-KO-screen_results.csv").open(encoding="utf-8-sig", newline="") as f:
        sr = list(csv.DictReader(f))
    survival_names = Counter(r["Gene"].strip() for r in sr)
    survival = {r["Gene"].strip(): r for r in sr if survival_names[r["Gene"].strip()] == 1 and not r["Gene"].strip().startswith("hsa-") and all(number(r[k]) for k in ["sgRNA number", "Cytokine survival beta", "Cytokine survival fdr"]) and float(r["sgRNA number"]) >= 3}
    sm = symbol_map("D1_survival_eligible", survival)
    book = openpyxl.load_workbook(RAW / "41588_2022_1261_MOESM4_ESM.xlsx", read_only=True, data_only=True)
    hits = {str(r[0]).strip(): r[2] for r in book["Table S1-Screen hits "].iter_rows(min_row=2, values_only=True) if r[0]}
    hm = symbol_map("D2_S1_published_hits", hits)
    guides = defaultdict(lambda: defaultdict(dict))
    for r in book["Table S8-Screen_all sgRNA"].iter_rows(min_row=2, values_only=True):
        match = re.fullmatch(r"(.*)_r([01])", str(r[0]).strip())
        if not match or not r[1] or not all(number(r[i]) for i in [2, 3, 6, 13]): continue
        key, rep = match.groups(); gene = str(r[1]).strip()
        assert key not in guides[gene][rep]
        guides[gene][rep][key] = (float(r[2]), float(r[3]))
    gm = symbol_map("D2_S8_measured_genes", set(guides) - {"EGFP", "LacZ", "luciferase"})
    covered = {}
    for gene, reps in guides.items():
        both = set(reps["0"]) & set(reps["1"])
        covered[gene] = sum(all(sum(reps[rep][key]) >= 20 for rep in ("0", "1")) for key in both)
    erows, ebatch = {}, []
    for i, r in enumerate(book["Table S5 - RNASeq"].iter_rows(min_row=2, max_col=3, values_only=True), 2):
        gid, symbol = str(r[0]).strip() if r[0] else "", str(r[1]).strip() if r[1] else ""
        # Version stripping is explicit and recorded; no other identifier edits.
        normalized = re.sub(r"\.\d+$", "", gid)
        targets = ensembl.get(normalized, set())
        hid = next(iter(targets)) if len(targets) == 1 else ""
        symbol_targets, symbol_method = resolve(symbol)
        reason = ""
        if not hid: reason = "ambiguous_Ensembl_id" if targets else "unmapped_Ensembl_id"
        elif len(symbol_targets) == 1 and hid not in symbol_targets: reason = "stable_ID_and_unique_symbol_disagree"
        elif not number(r[2]): reason = "nonfinite_baseMean"
        rec = {"source": "D2_S5_expression", "source_symbol": symbol, "source_GeneID": gid, "source_row": i, "mapping_method": "source_Ensembl_GeneID" + ("_version_removed" if gid != normalized else ""), "candidate_HGNC_ids": ";".join(sorted(targets)), "HGNC_id": hid, "approved_symbol": by_id[hid]["symbol"] if hid else "", "included_mapping": not bool(reason), "exclusion_reason": reason}
        ebatch.append(rec); erows[i] = (symbol, gid, r[2])
    book.close()
    ec = Counter(r["HGNC_id"] for r in ebatch if r["included_mapping"])
    for r in ebatch:
        if r["included_mapping"] and ec[r["HGNC_id"]] > 1:
            r["included_mapping"] = False; r["exclusion_reason"] = "multiple_expression_rows_to_one_HGNC_id"
    ledger.extend(ebatch)
    ei = {r["HGNC_id"]: r for r in ebatch if r["included_mapping"]}
    si = {h: g for g, h in sm.items()}
    gi = {h: g for g, h in gm.items()}
    hi = {h: g for g, h in hm.items()}
    joined = []
    for h in sorted(set(si) & set(gi) & set(ei)):
        sg, ig, er = si[h], gi[h], ei[h]
        expression_symbol, geneid, mean = erows[er["source_row"]]
        if covered[ig] < 3 or float(mean) < 10: continue
        s = survival[sg]; hit = hi.get(h, "")
        joined.append({"HGNC_id": h, "approved_symbol": by_id[h]["symbol"], "survival_source_symbol": sg, "insulin_source_symbol": ig,
            "expression_source_symbol": expression_symbol, "expression_GeneID": geneid, "expression_source_row": er["source_row"], "published_hit_source_symbol": hit,
            "baseMean": float(mean), "survival_guides": float(s["sgRNA number"]), "covered_insulin_guides": covered[ig],
            "A": float(s["Cytokine survival beta"]) > 0 and float(s["Cytokine survival fdr"]) < .1,
            "L": hits.get(hit) == "Decrease", "published_content_label": hits.get(hit, "Not a published hit")})
    frame = pd.DataFrame(joined)
    q, boundaries = pd.qcut(np.log2(frame.baseMean+1), 5, labels=False, retbins=True, duplicates="drop")
    frame["stratum"] = (q.astype(int)+1).astype(str) + "|" + np.where(frame.survival_guides >= 6, "6+", "3-5") + "|" + np.where(frame.covered_insulin_guides >= 4, "4+", "3")
    strata = []
    for e, s, g in itertools.product(range(1,len(boundaries)), ["3-5", "6+"], ["3", "4+"]):
        label = f"{e}|{s}|{g}"; subset = frame[frame.stratum == label]
        strata.append({"stratum": label, "n": len(subset), "m": int(subset.A.sum()), "k": int(subset.L.sum()), "a": int((subset.A & subset.L).sum())})
    folder = ROOT / "results/identifier_sensitivity"; folder.mkdir(exist_ok=True)
    frame.to_csv(folder / "stable_id_primary_genes.csv", index=False)
    write_csv(folder / "stable_id_primary_strata.csv", strata)
    write_csv(folder / "mapping_ledger.csv", ledger)
    n, a, l, overlap = len(frame), int(frame.A.sum()), int(frame.L.sum()), int((frame.A & frame.L).sum())
    result = {"design_status": "Post-review bounded identifier sensitivity; not the frozen exact-symbol primary", "HGNC_sha256": HGNC_HASH,
        "mapping_rules": "Exact approved symbol first; otherwise unique previous/alias symbol; reject within-source many-to-one collisions. Expression uses unique source Ensembl GeneID to HGNC ID, rejects disagreement with a uniquely mapped symbol and multiple expression rows per HGNC ID. Missing or ambiguous mappings are excluded. No inferred many-to-many matches.",
        "n": n, "A": a, "L": l, "overlap": overlap, "table_2x2": [[overlap,a-overlap],[l-overlap,n-a-l+overlap]],
        "log2_baseMean_plus1_quintile_boundaries": boundaries.tolist(), "conditional": exact_conditional(strata),
        "overlap_rows": frame[frame.A & frame.L].to_dict("records"),
        "mapping_dispositions": {source: dict(Counter(r["exclusion_reason"] or r["mapping_method"] for r in ledger if r["source"] == source)) for source in sorted({r["source"] for r in ledger})},
        "NUPL1_mapping": [r for r in ledger if r["source_symbol"] == "NUPL1" or r["HGNC_id"] in resolve("NUPL1")[0]]}
    original = ROOT / "results/primary_genes.csv"
    if original.exists():
        with original.open(encoding="utf-8", newline="") as f: exact_symbols = {r["gene"] for r in csv.DictReader(f)}
        mapped_symbols = set(frame.survival_source_symbol)
        result["comparison_to_exact_symbol_primary"] = {"exact_symbol_n": len(exact_symbols), "added_survival_source_symbols": sorted(mapped_symbols-exact_symbols), "removed_survival_source_symbols": sorted(exact_symbols-mapped_symbols)}
    result["source_hashes"] = {str(p.relative_to(ROOT)).replace("\\","/"): hashlib.sha256(p.read_bytes()).hexdigest() for p in [HGNC, RAW/"Supplementary_Data_09.CRISPR-KO-screen_results.csv", RAW/"41588_2022_1261_MOESM4_ESM.xlsx", Path(__file__)]}
    result["output_hashes"] = {str(p.relative_to(ROOT)).replace("\\","/"): hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.glob("*.csv")}
    (ROOT / "evidence/identifier_mapping_sensitivity.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps({k:result[k] for k in ["n","A","L","overlap","conditional","mapping_dispositions"]}, indent=2))

if __name__ == "__main__": main()
