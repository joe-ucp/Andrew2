"""Generate synchronized DOCX/TeX from a single versioned scientific content source."""
import argparse,json,re,shutil,subprocess,tempfile
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from prepare_manuscript import prepare
ROOT=Path(__file__).resolve().parents[1];MAN=ROOT/'manuscript'

def math_run(text,normal=False):
    run=OxmlElement('m:r')
    if normal:
        pr=OxmlElement('m:rPr');literal=OxmlElement('m:nor');literal.set(qn('m:val'),'1');pr.append(literal);run.append(pr)
    value=OxmlElement('m:t');value.text=text
    if normal:value.set(qn('xml:space'),'preserve')
    run.append(value);return run

def math_sub(base,subscript):
    node=OxmlElement('m:sSub')
    e=OxmlElement('m:e');e.append(math_run(base));node.append(e)
    s=OxmlElement('m:sub');s.append(math_run(subscript));node.append(s);return node

def binomial(top,bottom):
    d=OxmlElement('m:d');pr=OxmlElement('m:dPr')
    for name,value in [('begChr','('),('endChr',')')]:
        el=OxmlElement('m:'+name);el.set(qn('m:val'),value);pr.append(el)
    d.append(pr);e=OxmlElement('m:e');f=OxmlElement('m:f');fp=OxmlElement('m:fPr')
    typ=OxmlElement('m:type');typ.set(qn('m:val'),'noBar');fp.append(typ);f.append(fp)
    for name,items in [('num',top),('den',bottom)]:
        node=OxmlElement('m:'+name)
        for item in items:node.append(item)
        f.append(node)
    e.append(f);d.append(e);return d

def add_conditional_equation(doc):
    p=doc.add_paragraph();p.alignment=1
    para=OxmlElement('m:oMathPara');eq=OxmlElement('m:oMath')
    # Keep the conditioning bar literal and upright across OMML importers.
    items=[math_run('Pr('),math_sub('X','s'),math_run(' = x'),math_run(' | ',normal=True),math_sub('n','s'),math_run(', '),math_sub('m','s'),math_run(', '),math_sub('k','s'),math_run(', θ) ∝ '),binomial([math_sub('k','s')],[math_run('x')]),binomial([math_sub('n','s'),math_run(' − '),math_sub('k','s')],[math_sub('m','s'),math_run(' − x')])]
    sup=OxmlElement('m:sSup');e=OxmlElement('m:e');e.append(math_run('θ'));sup.append(e);power=OxmlElement('m:sup');power.append(math_run('x'));sup.append(power);items.append(sup)
    for item in items:eq.append(item)
    para.append(eq);p._p.append(para)

def tex(s):
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(mapping.get(c,c) for c in str(s)).replace('--','-{}-')

def prose_tex(s):
    parts=re.split(r'(https?://[^\s]+)',s)
    output=[]
    for part in parts:
        if part.startswith(('http://','https://')):
            url=part.rstrip('.,;');trail=part[len(url):]
            output.append(r'\url{'+url+'}'+tex(trail))
        else:output.append(tex(part))
    return ''.join(output)

def docx_table(doc,item):
    data=[item['headers']]+item['rows'];table=doc.add_table(rows=0,cols=len(data[0]));table.autofit=False
    widths=item.get('widths',[6.5/len(data[0])]*len(data[0]))
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    for idx,row in enumerate(data):
        cells=table.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.text=str(value)
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'40' if edge in ['top','bottom'] else '70');el.set(qn('w:type'),'dxa');mar.append(el)
            tcPr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEF2' if idx==0 else ('F6F7F8' if idx%2==0 else 'FFFFFF'));tcPr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1
                for run in p.runs:run.font.size=Pt(item.get('font_size',8.5));run.bold=(idx==0)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');table.rows[-1]._tr.get_or_add_trPr().append(repeat)
        cant=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(cant)
    doc.add_paragraph('')

def build():
    content=prepare()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.3)
    doc.core_properties.author='UCP Technology LLC';doc.core_properties.title=content['title'];doc.core_properties.subject='Reproducible analysis of beta cell CRISPR screens';doc.core_properties.keywords='beta cell; CRISPR; insulin content; cytokines'
    for style in doc.styles:
        if style.type==1:
            style.font.color.rgb=RGBColor(0,0,0)
            if style.element.pPr is not None:
                for b in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(b)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(10.5);normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
    for name,size in [('Title',19),('Heading 1',13),('Heading 2',11)]:
        doc.styles[name].font.name='Times New Roman';doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        rf=doc.styles[name].element.rPr.rFonts
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            if qn('w:'+attr) in rf.attrib:del rf.attrib[qn('w:'+attr)]
    hp=sec.header.paragraphs[0];hp.text='UCP Technology LLC  |  Beta cell screen comparison';hp.style='Normal'
    hp.runs[0].font.size=Pt(8)
    fp=sec.footer.paragraphs[0];fp.alignment=2
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
    t=[r'\documentclass[10pt,letterpaper]{article}',r'\usepackage[margin=0.8in]{geometry}',r'\usepackage[T1]{fontenc}',r'\usepackage{lmodern}',r'\usepackage{graphicx,booktabs,array,longtable,amsmath,xurl}',r'\usepackage[hidelinks]{hyperref}',r'\usepackage{fancyhdr}',r'\pagestyle{fancy}\fancyhf{}\fancyhead[L]{\footnotesize UCP Technology LLC | Beta cell screen comparison}\fancyfoot[R]{\thepage}',r'\renewcommand{\headrulewidth}{0pt}',r'\setlength{\headheight}{13pt}\setlength{\parindent}{0pt}\setlength{\parskip}{6pt}',r'\renewcommand{\arraystretch}{1.18}\setlength{\emergencystretch}{2em}',r'\hypersetup{pdftitle={'+tex(content['title'])+r'},pdfauthor={UCP Technology LLC}}',r'\begin{document}']
    for pageidx,page in enumerate(content['pages']):
        page_start=len(doc.paragraphs)
        if pageidx:t.append(r'\clearpage')
        for item in page:
            kind=item['type']
            if kind=='title':
                doc.add_paragraph(item['text'],'Title');t.append(r'{\LARGE\bfseries\raggedright '+tex(item['text'])+r'\par}')
            elif kind=='byline':
                p=doc.add_paragraph(item['text']);p.paragraph_format.space_after=Pt(12);t.append(r'{\large '+tex(item['text']).replace('\n',r'\\')+r'\par}')
            elif kind=='heading':
                level=item.get('level',1);doc.add_heading(item['text'],level);t.append((r'\section*{' if level==1 else r'\subsection*{')+tex(item['text'])+'}')
            elif kind in ('paragraph','reference'):
                p=doc.add_paragraph(item['text'])
                if kind=='reference':
                    p.paragraph_format.space_after=Pt(3)
                    for run in p.runs:run.font.size=Pt(9)
                    t.append(r'{\small '+prose_tex(item['text'])+'}\n')
                else:t.append(prose_tex(item['text'])+'\n')
            elif kind=='equation':
                add_conditional_equation(doc)
                t.append(r'\['+item['latex']+r'\]')
            elif kind=='table':
                p=doc.add_paragraph(item['caption']);p.runs[0].bold=True
                docx_table(doc,item)
                widths=item.get('widths',[6.5/len(item['headers'])]*len(item['headers']))
                spec=''.join(r'>{\raggedright\arraybackslash}p{'+f'{w*.91:.2f}in'+'}' for w in widths)
                t.extend([r'{\small\textbf{'+tex(item['caption'])+r'}\par}',r'{\small\begin{longtable}{'+spec+r'}\toprule',' & '.join(r'\textbf{'+tex(v)+'}' for v in item['headers'])+r'\\\midrule\endhead'])
                t.extend(' & '.join(tex(v) for v in row)+r'\\' for row in item['rows'])
                t.append(r'\bottomrule\end{longtable}}')
            elif kind=='figure':
                pic=doc.add_picture(str(MAN/item['path']),width=Inches(item.get('width',6.5)))
                pic._inline.docPr.set('descr',item['alt'])
                pic._inline.docPr.set('title',item.get('title',item['caption'].split('.')[0]))
                p=doc.add_paragraph(item['caption']);p.paragraph_format.space_after=Pt(8)
                for run in p.runs:run.font.size=Pt(9)
                t.append(r'\begin{center}\includegraphics[width='+str(item.get('width',6.5))+r'in]{'+item['path']+r'}\end{center}')
                t.append(r'{\small '+tex(item['caption'])+r'\par}')
        if pageidx:doc.paragraphs[page_start].paragraph_format.page_break_before=True
    t.append(r'\end{document}')
    doc.save(MAN/'paper.docx');(MAN/'paper.tex').write_text('\n'.join(t)+'\n',encoding='utf-8')

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--compile',action='store_true');args=ap.parse_args();build()
    if args.compile:
        # Keep compiler scratch files outside the research repository.
        with tempfile.TemporaryDirectory(prefix='beta-cell-manuscript-') as temporary:
            for _ in range(2):
                subprocess.run(['pdflatex','-interaction=nonstopmode','-halt-on-error',
                                '-output-directory='+temporary,'paper.tex'],
                               cwd=MAN,check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
            shutil.copyfile(Path(temporary)/'paper.pdf',MAN/'paper.pdf')
    print('Built synchronized manuscript/paper.docx and paper.tex'+(' and paper.pdf' if args.compile else ''))
