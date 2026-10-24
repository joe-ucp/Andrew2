"""Independent stdlib/openpyxl reconstruction; never imports the main analysis.

Default mode audits source marginals only. --joint also reconstructs and compares
the cross-screen analyses. Primary source joins are exact trimmed symbols.
"""
from __future__ import annotations

import argparse
import bisect
import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from statistics import median
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
EVIDENCE = ROOT / "evidence"
EXPECTED = {
    "Supplementary_Data_09.CRISPR-KO-screen_results.csv": "d14accf5fa62320c114741de936c5f3ca05754babd0257670a55542a7c1187bc",
    "41588_2022_1261_MOESM4_ESM.xlsx": "a8f47504edd26d999cb861f146f6fd6aa2ee2ec8368a01e29c808a10238a7690",
    "mmc6.xlsx": "e8f7c71f8f47b94bef33f1f90b487a7bcbf32f4328d61460c60080cd4c7abb8c",
    "16_CRISPR_figure.ipynb": "4b9da99ec36eb2d903a2b7fb987ceed24a7f56c9910da4860566171ee0360956",
}

def finite(value):
    try:
        return math.isfinite(float(value))
    except (ValueError, TypeError):
        return False

def clean(value):
    return str(value).strip() if value is not None else ""

def read_sources():
    for file, expected in EXPECTED.items():
        assert hashlib.sha256((RAW / file).read_bytes()).hexdigest() == expected, file
    with (RAW / "Supplementary_Data_09.CRISPR-KO-screen_results.csv").open(encoding="utf-8-sig", newline="") as f:
        survival_rows = list(csv.DictReader(f))
    symbols = Counter(clean(r["Gene"]) for r in survival_rows)
    survival = {}
    for row in survival_rows:
        g = clean(row["Gene"])
        if g and symbols[g] == 1 and not g.startswith("hsa-") and all(finite(row[k]) for k in ("sgRNA number", "Cytokine survival beta", "Cytokine survival fdr")) and float(row["sgRNA number"]) >= 3:
            survival[g] = {"beta": float(row["Cytokine survival beta"]), "fdr": float(row["Cytokine survival fdr"]), "nguides": int(row["sgRNA number"])}
    wb = openpyxl.load_workbook(RAW / "41588_2022_1261_MOESM4_ESM.xlsx", data_only=True, read_only=True)
    hits = {}
    for row in wb["Table S1-Screen hits "].iter_rows(min_row=2, max_col=3, values_only=True):
        g = clean(row[0])
        if not g:
            continue
        assert g not in hits, f"Duplicate S1 symbol: {g}"
        hits[g] = clean(row[2])
    expressions = defaultdict(list)
    for row in wb["Table S5 - RNASeq"].iter_rows(min_row=2, max_col=3, values_only=True):
        g = clean(row[1])
        if g:
            expressions[g].append(row[2])
    expression = {g: float(vals[0]) for g, vals in expressions.items() if len(vals) == 1 and finite(vals[0])}
    expression_duplicates = sorted(g for g, vals in expressions.items() if len(vals) > 1)
    entries = defaultdict(lambda: defaultdict(list))
    gene_labels = set()
    replicas = Counter()
    invalid_guide_rows = []
    rows_count = 0
    for i, row in enumerate(wb["Table S8-Screen_all sgRNA"].iter_rows(min_row=2, values_only=True), 2):
        if row[0] is None:
            continue
        rows_count += 1
        g = clean(row[1])
        gene_labels.add(g)
        match = re.fullmatch(r"(.*)_r([01])", clean(row[0]))
        if not match:
            invalid_guide_rows.append(i)
            continue
        key, replicate = match.group(1), match.group(2)
        replicas[replicate] += 1
        values = [row[2], row[3], row[6], row[13]]
        valid = bool(g) and all(finite(v) for v in values)
        entry = {"gene": g, "row": i, "valid": valid}
        if valid:
            entry.update(control=float(values[0]), treatment=float(values[1]), lfc=float(values[2]), fdr=float(values[3]))
        else:
            invalid_guide_rows.append(i)
        entries[key][replicate].append(entry)
    wb.close()
    coverage = Counter()
    support = defaultdict(lambda: {"0": {1: set(), -1: set()}, "1": {1: set(), -1: set()}})
    lfcs = defaultdict(lambda: {"0": [], "1": []})
    duplicates = []
    mismatched_genes = []
    for key, reps in entries.items():
        for rep, data in reps.items():
            if len(data) != 1:
                duplicates.append([key, rep, len(data)])
                continue
            rec = data[0]
            if rec["valid"]:
                lfcs[rec["gene"]][rep].append(rec["lfc"])
                if rec["fdr"] < .1 and rec["lfc"] != 0:
                    support[rec["gene"]][rep][1 if rec["lfc"] > 0 else -1].add(key)
        if set(reps) != {"0", "1"} or any(len(data) != 1 for data in reps.values()):
            continue
        r0, r1 = reps["0"][0], reps["1"][0]
        if r0["gene"] != r1["gene"]:
            mismatched_genes.append(key)
            continue
        if all(r["valid"] and r["control"] + r["treatment"] >= 20 for r in (r0, r1)):
            coverage[r0["gene"]] += 1
    covered = {g for g, count in coverage.items() if count >= 3 and g not in {"EGFP", "LacZ", "luciferase"}}
    audit_failures = []
    for g, label in hits.items():
        assert label in {"Decrease", "Increase"}, (g, label)
        sign = 1 if label == "Decrease" else -1
        if any(len(support[g][rep][sign]) < 2 for rep in ("0", "1")):
            audit_failures.append(g)
    U0 = set(survival) & covered
    U = {g for g in U0 if g in expression and expression[g] >= 10 and g not in audit_failures}
    source_audit = {
        "scope": "Source margins and eligibility only; no cross-screen overlap or association computed in this object.",
        "survival_rows": len(survival_rows), "survival_eligible": len(survival),
        "survival_positive_fdr_0_10": sum(r["beta"] > 0 and r["fdr"] < .1 for r in survival.values()),
        "survival_negative_fdr_0_10": sum(r["beta"] < 0 and r["fdr"] < .1 for r in survival.values()),
        "survival_duplicate_names": [g for g, n in symbols.items() if n > 1],
        "insulin_hit_labels": dict(Counter(hits.values())), "insulin_guide_rows": rows_count,
        "insulin_rows_per_replicate": dict(replicas), "insulin_unique_gene_labels": len(gene_labels),
        "insulin_coverage_eligible": len(covered), "guide_duplicate_keys": duplicates,
        "guide_gene_mismatches": mismatched_genes, "invalid_guide_rows": invalid_guide_rows,
        "expression_duplicate_names_count": len(expression_duplicates), "expression_duplicate_names": expression_duplicates,
        "directional_audit_failures": audit_failures,
        "directional_audit_failure_details": {g: {"label": hits[g], "positive_r0": len(support[g]["0"][1]), "positive_r1": len(support[g]["1"][1]), "negative_r0": len(support[g]["0"][-1]), "negative_r1": len(support[g]["1"][-1]), "baseMean": expression.get(g)} for g in audit_failures},
        "U0_before_direction_audit": len(U0), "U_primary": len(U),
        "primary_survival_favoring_margin": sum(survival[g]["beta"] > 0 and survival[g]["fdr"] < .1 for g in U),
        "primary_decrease_margin": sum(hits.get(g) == "Decrease" for g in U),
        "SOCS1_survival": survival.get("SOCS1"),
        "INS_content": {rep: {"median_lfc": median(lfcs["INS"][rep]), "positive_fdr_lt_0_10_guides": len(support["INS"][rep][1])} for rep in ("0", "1")},
    }
    d3book = openpyxl.load_workbook(RAW / "mmc6.xlsx", data_only=True, read_only=True)
    d3rows = list(d3book["KO screening results"].iter_rows(min_row=3, max_col=6, values_only=True))
    d3book.close()
    d1bygene = {clean(r["Gene"]): r for r in survival_rows}
    mismatches = []
    keys = ("sgRNA number", "Cytokine survival beta", "Cytokine survival z", "Cytokine survival p-value", "Cytokine survival fdr")
    for row in d3rows:
        g = clean(row[0])
        if g not in d1bygene:
            mismatches.append([g, "missing_from_D1"])
        else:
            for key, value in zip(keys, row[1:]):
                if not math.isclose(float(value), float(d1bygene[g][key]), rel_tol=1e-12, abs_tol=1e-15):
                    mismatches.append([g, key, value, d1bygene[g][key]])
    rounded = [r for r in mismatches if len(r) == 4 and "E" in r[3].upper() and float(format(float(r[2]), ".2E")) == float(r[3])]
    unaccounted = [r for r in mismatches if r not in rounded]
    source_audit["D3_crosscheck"] = {
        "rows": len(d3rows), "D1_rows": len(survival_rows),
        "same_gene_membership": {clean(r[0]) for r in d3rows} == set(d1bygene),
        "numerical_tolerance": "relative 1e-12; absolute 1e-15",
        "numerical_differences_count": len(mismatches),
        "scientific_notation_rounding_count": len(rounded),
        "rounding_rule": "D1 number equals D3 formatted to two decimal places in scientific notation; all full values and source strings retained below.",
        "unexplained_differences": unaccounted,
        "any_FDR_difference": any(r[1] == "Cytokine survival fdr" for r in mismatches),
        "any_beta_sign_change": any(r[1] == "Cytokine survival beta" and (float(r[2]) > 0) != (float(r[3]) > 0) for r in mismatches),
        "differences": mismatches,
    }
    notebook = json.loads((RAW / "16_CRISPR_figure.ipynb").read_text(encoding="utf-8"))
    source_audit["notebook_filter_and_table"] = [{"source": c.get("source"), "outputs": c.get("outputs")} for c in notebook["cells"] if "sgRNA >2" in "".join(c.get("source", [])) or "table(res$cytokine_survival" in "".join(c.get("source", []))]
    source_audit["denominator_discrepancy"] = {"article": 18703, "pinned_data": len(survival), "notebook_displayed_counts_sum": 9043 + 8804 + 440 + 427, "action": "Retain explicit 18,714 eligible rows; do not delete 11 unspecified genes."}
    return source_audit, survival, hits, expression, coverage, support, U0, U, lfcs

def independently_construct_primary(data, universe=None, expression_strata=True, survival_fdr=.1, negative_exposure=False, min_directional=2, fixed_edges=None):
    """Use stdlib quantile interpolation and label reconstruction from raw inputs."""
    _, survival, hits, expression, coverage, support, U0, U, lfcs = data
    U = set(U if universe is None else universe)
    values = sorted(math.log2(expression[g] + 1) for g in U) if expression_strata else []
    def quantile(p):
        h = (len(values) - 1) * p
        lower = math.floor(h)
        upper = math.ceil(h)
        return values[lower] + (values[upper] - values[lower]) * (h - lower)
    edges = (list(fixed_edges) if fixed_edges is not None else sorted(set(quantile(p) for p in [0, .2, .4, .6, .8, 1]))) if expression_strata else [0, 1]
    boundaries = edges[1:-1]
    rows = []
    counts = {(q, s, i): [0, 0, 0, 0] for q in range(len(edges) - 1) for s in (0, 1) for i in (0, 1)}
    for g in sorted(U):
        a = int((survival[g]["beta"] < 0 if negative_exposure else survival[g]["beta"] > 0) and survival[g]["fdr"] < survival_fdr)
        l = int(hits.get(g) == "Decrease" and all(len(support[g][rep][1]) >= min_directional for rep in ("0", "1")))
        q = bisect.bisect_left(boundaries, math.log2(expression[g] + 1)) if expression_strata else 0
        s = int(survival[g]["nguides"] >= 6)
        i = int(coverage[g] >= 4)
        counts[q, s, i][0] += 1
        counts[q, s, i][1] += a
        counts[q, s, i][2] += l
        counts[q, s, i][3] += a * l
        q_label = q + 1 if expression_strata else 0
        rows.append({"gene": g, "A": a, "L": l, "expression_quintile": q_label, "survival_coverage_high": s, "insulin_coverage_high": i,
                     "stratum": f"{q_label}|{'6+' if s else '3-5'}|{'4+' if i else '3'}",
                     "survival_beta": survival[g]["beta"], "survival_fdr": survival[g]["fdr"], "survival_guides": survival[g]["nguides"],
                     "insulin_covered_guides": coverage[g], "baseMean": expression.get(g), "published_content_label": hits.get(g, ""),
                     "r0_median_lfc": median(lfcs[g]["0"]), "r1_median_lfc": median(lfcs[g]["1"])})
    strata = [{"expression_quintile": q + 1 if expression_strata else 0, "stratum": f"{q + 1 if expression_strata else 0}|{'6+' if s else '3-5'}|{'4+' if i else '3'}", "survival_coverage_high": s, "insulin_coverage_high": i, "n": vals[0], "m": vals[1], "k": vals[2], "a": vals[3]} for (q, s, i), vals in sorted(counts.items())]
    return edges, rows, strata

def independent_statistics(strata):
    """SciPy distribution implementation, separate from main log-weight routine."""
    import numpy as np
    from scipy.optimize import brentq
    from scipy.stats import nchypergeom_fisher, fisher_exact
    from scipy.stats.contingency import odds_ratio
    specifications = []
    n_total = m_total = k_total = a_total = 0
    for s in strata:
        n, m, k, a = (s[v] for v in ("n", "m", "k", "a"))
        lo, hi = max(0, m + k - n), min(m, k)
        specifications.append((n, m, k, lo, hi))
        n_total += n
        m_total += m
        k_total += k
        a_total += a
    t_min = sum(s[3] for s in specifications)
    t_max = sum(s[4] for s in specifications)
    observed_index = a_total - t_min
    def distribution(log_theta):
        pmf = np.array([1.0])
        for n, m, k, lo, hi in specifications:
            if lo == hi:
                p = np.array([1.0])
            else:
                p = nchypergeom_fisher.pmf(np.arange(lo, hi + 1), n, k, m, math.exp(log_theta))
                p = p / p.sum()
            pmf = np.convolve(pmf, p)
        return pmf / pmf.sum()
    def mean(log_theta):
        return sum(lo if lo == hi else float(nchypergeom_fisher.mean(n, k, m, math.exp(log_theta))) for n, m, k, lo, hi in specifications)
    def solve(fn):
        left, right = -1., 1.
        while fn(left) * fn(right) > 0:
            left *= 2
            right *= 2
            if right > 128:
                raise RuntimeError("Independent inversion could not bracket root")
        return math.exp(brentq(fn, left, right, xtol=1e-11, rtol=1e-12))
    if t_min == t_max:
        cmle, lower, upper = None, 0., "Infinity"
    else:
        cmle = 0. if a_total == t_min else "Infinity" if a_total == t_max else solve(lambda x: mean(x) - a_total)
        lower = 0. if a_total == t_min else solve(lambda x: float(distribution(x)[observed_index:].sum()) - .025)
        upper = "Infinity" if a_total == t_max else solve(lambda x: float(distribution(x)[:observed_index + 1].sum()) - .025)
    upper_p = float(distribution(0.)[observed_index:].sum())
    table = [[a_total, m_total - a_total], [k_total - a_total, n_total - m_total - k_total + a_total]]
    fisher = odds_ratio(table, kind="conditional")
    interval = fisher.confidence_interval(confidence_level=.95)
    return {"stratified_cmle": cmle, "stratified_ci95": [lower, upper], "stratified_upper_p": upper_p,
            "unstratified_table_A_by_L": table, "fisher_conditional_or": float(fisher.statistic),
            "fisher_exact_ci95": [float(interval.low), float(interval.high)], "fisher_upper_p": float(fisher_exact(table, alternative="greater").pvalue),
            "support_T": [t_min, t_max]}

def joint_verification(data):
    inputs = json.loads((EVIDENCE / "analysis_inputs.json").read_text(encoding="utf-8"))
    for item in inputs["files"]:
        assert hashlib.sha256((ROOT / item["path"]).read_bytes()).hexdigest() == item["sha256"], f"Analysis input changed: {item['path']}"
    _, survival, hits, expression, coverage, support, U0, U, lfcs = data
    failures2 = set(data[0]["directional_audit_failures"])
    failures3 = {g for g, label in hits.items() if any(len(support[g][r][1 if label == "Decrease" else -1]) < 3 for r in ("0", "1"))}
    edges, primary_rows, primary_strata = independently_construct_primary(data)
    analyses = {"primary": (edges, primary_rows, primary_strata)}
    analyses["no_expression"] = independently_construct_primary(data, universe=U0-failures2, expression_strata=False)
    for threshold in (1, 100):
        universe = {g for g in U0-failures2 if g in expression and expression[g] >= threshold}
        analyses[f"expression_ge{threshold}"] = independently_construct_primary(data, universe=universe)
    analyses["survival_fdr005"] = independently_construct_primary(data, survival_fdr=.05)
    analyses["three_directional_guides"] = independently_construct_primary(data, universe=U-failures3, min_directional=3)
    analyses["survival_harming"] = independently_construct_primary(data, negative_exposure=True)
    hit_subset = {g for g in U if survival[g]["fdr"] < .1 and survival[g]["beta"] != 0}
    analyses["positive_vs_negative"] = independently_construct_primary(data, universe=hit_subset, fixed_edges=edges)
    analyses["three_guides_fixed_universe"] = independently_construct_primary(data, min_directional=3)
    output = EVIDENCE / "independent"
    output.mkdir(exist_ok=True)
    checks = {"implementation": "Standalone csv/openpyxl and stdlib source construction; manual linear quantiles/bisect; SciPy noncentral hypergeometric PMFs and independent convolution/root inversion. No import of analyze.py or exact_stats.py.", "analysis_input_hashes_verified": True, "analyses": {}, "candidate_checks": {}}
    main_results = {r["name"]: r for r in json.loads((ROOT / "results/results.json").read_text())["analyses"]}
    for name, (bins, rows, strata) in analyses.items():
        with (output / f"{name}_genes.csv").open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]))
            writer.writeheader(); writer.writerows(rows)
        with (output / f"{name}_strata.csv").open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(strata[0]))
            writer.writeheader(); writer.writerows(strata)
        with (ROOT / f"results/{name}_genes.csv").open(encoding="utf-8", newline="") as f:
            reference_rows = {r["gene"]: r for r in csv.DictReader(f)}
        with (ROOT / f"results/{name}_strata.csv").open(encoding="utf-8", newline="") as f:
            reference_strata = {r["stratum"]: r for r in csv.DictReader(f)}
        independent_rows = {r["gene"]: r for r in rows}
        differences = []
        for gene in sorted(set(reference_rows) | set(independent_rows)):
            if gene not in reference_rows or gene not in independent_rows:
                differences.append([gene, "membership"])
                continue
            a, b = independent_rows[gene], reference_rows[gene]
            for field in ("A", "L"):
                if bool(a[field]) != (b[field] == "True"):
                    differences.append([gene, field, a[field], b[field]])
            if a["stratum"] != b["stratum"]:
                differences.append([gene, "stratum", a["stratum"], b["stratum"]])
            for field, reference_field in [("r0_median_lfc", "median_lfc_r0"), ("r1_median_lfc", "median_lfc_r1"), ("insulin_covered_guides", "covered_insulin_guides"), ("survival_beta", "survival_beta"), ("survival_fdr", "survival_fdr")]:
                if not math.isclose(a[field], float(b[reference_field]), rel_tol=1e-12, abs_tol=1e-12):
                    differences.append([gene, field, a[field], b[reference_field]])
        stratum_differences = []
        for s in strata:
            ref = reference_strata.get(s["stratum"])
            if ref is None or any(s[k] != int(ref[k]) for k in ("n", "m", "k", "a")):
                stratum_differences.append([s, ref])
        stats = independent_statistics(strata)
        reference_stats = main_results[name]
        stats_differences = {}
        comparisons = {
            "conditional_mle": (stats["stratified_cmle"], reference_stats["conditional_mle"]),
            "ci_lower": (stats["stratified_ci95"][0], reference_stats["ci_lower"]),
            "ci_upper": (stats["stratified_ci95"][1], reference_stats["ci_upper"]),
            "p_enrichment": (stats["stratified_upper_p"], reference_stats["p_enrichment"]),
            "fisher_conditional_mle": (stats["fisher_conditional_or"], reference_stats["unstratified"]["conditional_mle"]),
            "fisher_ci_lower": (stats["fisher_exact_ci95"][0], reference_stats["unstratified"]["ci_lower"]),
            "fisher_ci_upper": (stats["fisher_exact_ci95"][1], reference_stats["unstratified"]["ci_upper"]),
            "fisher_p_enrichment": (stats["fisher_upper_p"], reference_stats["unstratified"]["p_enrichment"]),
        }
        for key, (a, b) in comparisons.items():
            if not math.isclose(a, b, rel_tol=1e-7, abs_tol=1e-10):
                stats_differences[key] = [a, b]
        checks["analyses"][name] = {"gene_count": len(rows), "stratum_count": len(strata), "bins": bins, "gene_differences": differences, "stratum_differences": stratum_differences, "statistic_differences": stats_differences, "statistics": stats, "passed": not (differences or stratum_differences or stats_differences)}
        print(f"Independent {name}: N={len(rows)} table={stats['unstratified_table_A_by_L']} passed={checks['analyses'][name]['passed']}", flush=True)
    source_candidates = {g for g, r in survival.items() if r["beta"] > 0 and r["fdr"] < .1}
    with (ROOT / "results/candidate_table_427.csv").open(encoding="utf-8", newline="") as f:
        candidates = list(csv.DictReader(f))
    class_counts = Counter()
    candidate_differences = []
    for row in candidates:
        g = row["gene"]
        if g not in U:
            category = "Comparison unavailable or inconsistent"
        elif hits.get(g) == "Decrease":
            category = "Supported cross-screen conflict"
        elif hits.get(g) == "Increase":
            category = "Survival-favoring with higher-content evidence"
        else:
            category = "Content effect unresolved"
        class_counts[category] += 1
        if row["candidate_class"] != category:
            candidate_differences.append([g, "class", category, row["candidate_class"]])
        for key, source_key in (("survival_beta", "beta"), ("survival_fdr", "fdr"), ("survival_guides", "nguides")):
            if not math.isclose(float(row[key]), survival[g][source_key], rel_tol=1e-12, abs_tol=1e-12):
                candidate_differences.append([g, key, row[key], survival[g][source_key]])
        reasons = set()
        if g not in lfcs:
            reasons.add("no_exact_insulin_symbol")
        elif coverage[g] < 3:
            reasons.add("covered_insulin_guides_lt3")
        if g in data[0]["expression_duplicate_names"]:
            reasons.add("ambiguous_expression_symbol")
        elif g not in expression:
            reasons.add("no_exact_expression_symbol")
        elif expression[g] < 10:
            reasons.add("expression_lt10")
        if g in failures2:
            reasons.add("failed_direction_audit")
        if set(row["eligibility_reason"].split(";")) != (reasons or {"eligible"}):
            candidate_differences.append([g, "eligibility_reason", sorted(reasons), row["eligibility_reason"]])
        for rep in ("0", "1"):
            for direction, sign in (("positive", 1), ("negative", -1)):
                actual = len(support[g][rep][sign])
                if actual != int(row[f"{direction}_support_r{rep}"]):
                    candidate_differences.append([g, f"{direction}_support_r{rep}", actual, row[f"{direction}_support_r{rep}"]])
        if not row["external_functional_evidence_status"].startswith("Unknown;"):
            candidate_differences.append([g, "external_evidence_not_unknown"])
    symbols_equal = {r["gene"] for r in candidates} == source_candidates and len(candidates) == len(source_candidates)
    checks["candidate_checks"] = {"all_source_427_accounted_exactly_once": symbols_equal, "class_counts": dict(class_counts), "differences": candidate_differences, "passed": symbols_equal and not candidate_differences}
    diagnostic_differences = []
    def quantile(values, p):
        values = sorted(values)
        h = (len(values) - 1) * p
        a, b = math.floor(h), math.ceil(h)
        return values[a] + (values[b] - values[a]) * (h-a)
    with (ROOT / "results/threshold_free_stratum_diagnostics.csv").open(encoding="utf-8", newline="") as f:
        for d in csv.DictReader(f):
            rep = d["replicate"]
            groups = {a: [r[f"{rep}_median_lfc"] for r in primary_rows if r["stratum"] == d["stratum"] and r["A"] == a] for a in (0, 1)}
            for a, values in groups.items():
                if len(values) != int(d[f"n_A{a}"]):
                    diagnostic_differences.append([d["stratum"], rep, f"n_A{a}"])
                for key, p in (("q25", .25), ("median", .5), ("q75", .75)):
                    actual = quantile(values, p) if values else None
                    reference = d[f"{key}_A{a}"]
                    if (actual is None and reference != "") or (actual is not None and (reference == "" or not math.isclose(actual, float(reference), rel_tol=1e-12, abs_tol=1e-12))):
                        diagnostic_differences.append([d["stratum"], rep, f"{key}_A{a}", actual, reference])
            difference = median(groups[1]) - median(groups[0]) if groups[1] and groups[0] else None
            reference = d["median_difference_A1_minus_A0"]
            if (difference is None and reference != "") or (difference is not None and (reference == "" or not math.isclose(difference, float(reference), rel_tol=1e-12, abs_tol=1e-12))):
                diagnostic_differences.append([d["stratum"], rep, "median_difference", difference, reference])
    checks["threshold_free_diagnostic"] = {"differences": diagnostic_differences, "passed": not diagnostic_differences}
    checks["passed"] = all(v["passed"] for v in checks["analyses"].values()) and checks["candidate_checks"]["passed"] and checks["threshold_free_diagnostic"]["passed"]
    (EVIDENCE / "independent_verification.json").write_text(json.dumps(checks, indent=2), encoding="utf-8")
    assert checks["passed"], "Independent verification discrepancies; see evidence/independent_verification.json"
    print("Independent joint and candidate verification PASSED", flush=True)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--joint", action="store_true")
    args = ap.parse_args()
    data = read_sources()
    EVIDENCE.mkdir(exist_ok=True, parents=True)
    (EVIDENCE / "source_marginal_audit.json").write_text(json.dumps(data[0], indent=2), encoding="utf-8")
    if args.joint:
        joint_verification(data)
    else:
        print(json.dumps({k: v for k, v in data[0].items() if k not in {"expression_duplicate_names", "notebook_filter_and_table", "D3_crosscheck"}}, indent=2))

if __name__ == "__main__":
    main()
