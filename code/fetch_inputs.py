"""Fetch hash-pinned public inputs without constructing cross-screen associations."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import tempfile
import urllib.error
import urllib.request
import zipfile
from pathlib import Path
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
EVIDENCE = ROOT / "evidence"
SOURCES = [
    {"id": "D1", "file": "Supplementary_Data_09.CRISPR-KO-screen_results.csv", "sha256": "d14accf5fa62320c114741de936c5f3ca05754babd0257670a55542a7c1187bc", "url": "https://zenodo.org/records/7106728/files/Supplementary_Data_09.CRISPR-KO-screen_results.csv?download=1"},
    {"id": "D2", "file": "41588_2022_1261_MOESM4_ESM.xlsx", "sha256": "a8f47504edd26d999cb861f146f6fd6aa2ee2ec8368a01e29c808a10238a7690", "url": "https://media.springernature.com/full/springer-static/esm/art%3A10.1038%2Fs41588-022-01261-2/MediaObjects/41588_2022_1261_MOESM4_ESM.xlsx", "source_archive_url": "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC9839450/supplementaryFiles"},
    {"id": "D3", "file": "mmc6.xlsx", "sha256": "e8f7c71f8f47b94bef33f1f90b487a7bcbf32f4328d61460c60080cd4c7abb8c", "url": "https://ars.els-cdn.com/content/image/1-s2.0-S2666979X22001677-mmc6.xlsx", "source_archive_url": "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC9903835/supplementaryFiles"},
    {"id": "C1", "file": "16_CRISPR_figure.ipynb", "sha256": "4b9da99ec36eb2d903a2b7fb987ceed24a7f56c9910da4860566171ee0360956", "hash_origin": "Pinned from the specified immutable source commit.", "url": "https://raw.githubusercontent.com/Gaulton-Lab/BetaCells_cytokines_T1D/7d10581be640a2a8f4586c53a677b7c6cbbd7b1c/notebooks/16_CRISPR_figure.ipynb", "commit": "7d10581be640a2a8f4586c53a677b7c6cbbd7b1c"},
]

ROTTNER_ARCHIVE = {
    "id": "rottner_author_code_v1",
    "file": "Rottner_etal_2022-v1.0.0.zip",
    "sha256": "45868253fd8ede3575900d6eaa51f54d307270b901d8e4ae49ed1682887b6dbd",
    "url": "https://zenodo.org/api/records/7226348/files/TGDStanford/Rottner_etal_2022-v1.0.0.zip/content",
}


def digest(data):
    return hashlib.sha256(data).hexdigest()

def resolve_rottner_archive(cache_dir=None, offline=False):
    """Resolve verified author code outside the repository; never redistribute it."""
    cache = Path(cache_dir or os.environ.get("BETA_CELL_SOURCE_CACHE") or
                 Path.home() / ".cache" / "beta-cell-screen-comparison" / "sources").expanduser().resolve()
    path = cache / ROTTNER_ARCHIVE["file"]
    if cache.is_relative_to(ROOT.resolve()) or path.resolve().is_relative_to(ROOT.resolve()):
        raise ValueError("The author-code cache must be outside the repository; use --source-cache or BETA_CELL_SOURCE_CACHE.")
    expected = ROTTNER_ARCHIVE["sha256"]
    if path.exists():
        actual = digest(path.read_bytes())
        if actual != expected:
            raise ValueError(f"HASH MISMATCH for cached {path.name}: expected {expected}, got {actual}; move the invalid cache file aside before fetching again.")
        return path
    if offline:
        raise FileNotFoundError(f"Missing remote-only author archive at {path}. Run python code/fetch_inputs.py --source-cache \"{cache}\" once with network access, then retry --offline.")
    request = urllib.request.Request(ROTTNER_ARCHIVE["url"], headers={"User-Agent": "BetaCellScreenComparison/1.0"})
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = response.read()
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Could not retrieve the remote-only author archive from {ROTTNER_ARCHIVE['url']}. Retry with network access or supply a verified external cache via --source-cache.") from exc
    actual = digest(payload)
    if actual != expected:
        raise ValueError(f"HASH MISMATCH for downloaded {path.name}: expected {expected}, got {actual}; nothing was cached.")
    cache.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=cache, prefix=".rottner-", suffix=".part", delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(payload)
        temporary.replace(path)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return path


def retrieve(source, offline=False):
    path = RAW / source["file"]
    resolved_url = source["url"]
    archive_member = None
    downloaded = False
    if path.exists():
        data = path.read_bytes()
    else:
        if offline:
            raise FileNotFoundError(f"Missing pinned input: {path}")
        req = urllib.request.Request(source["url"], headers={"User-Agent": "BetaCellScreenComparison/1.0"})
        with urllib.request.urlopen(req, timeout=180) as response:
            payload = response.read()
            resolved_url = response.url
        if source.get("archive"):
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                names = [n for n in archive.namelist() if Path(n).name == source["file"]]
                if len(names) != 1:
                    raise ValueError(f"Expected one exact archive member for {source['file']}, found {names}")
                archive_member = names[0]
                data = archive.read(archive_member)
        else:
            data = payload
        downloaded = True
    actual = digest(data)
    expected = source.get("sha256")
    if expected and actual != expected:
        raise ValueError(f"HASH MISMATCH for {path.name}: expected {expected}, got {actual}; do not silently accept source change")
    if downloaded:
        path.write_bytes(data)
    return {**source, "sha256": actual, "bytes": len(data), "path": str(path.relative_to(ROOT)).replace("\\", "/"), "resolved_url": resolved_url, "archive_member": archive_member, "hash_matches_expected": actual == expected if expected else None}

def schema():
    import openpyxl
    result = {"scope": "Schema and separate source margins only; no cross-screen association computed."}
    with (RAW / SOURCES[0]["file"]).open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        result["D1"] = {"columns": reader.fieldnames, "data_rows": len(rows), "first_rows": rows[:2]}
    for source in SOURCES[1:3]:
        book = openpyxl.load_workbook(RAW / source["file"], read_only=True, data_only=True)
        result[source["id"]] = {"sheets": {}}
        for ws in book:
            rows = [{"row_1based": i, "values": list(row)} for i, row in enumerate(ws.iter_rows(max_row=12, values_only=True), 1) if any(v is not None for v in row)]
            result[source["id"]]["sheets"][ws.title] = {"max_row": ws.max_row, "max_column": ws.max_column, "first_rows": rows}
        book.close()
    (EVIDENCE / "source_schema.json").write_text(json.dumps(result, indent=2, default=str), encoding="utf-8")
    return result

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--offline", action="store_true", help="Verify supplied files and the external author-code cache without downloading")
    ap.add_argument("--source-cache", help="External cache directory for remote-only author code; defaults to BETA_CELL_SOURCE_CACHE or ~/.cache/beta-cell-screen-comparison/sources")
    ap.add_argument("--schema", action="store_true", help="Also inspect schemas without cross-screen joining")
    args = ap.parse_args()
    RAW.mkdir(parents=True, exist_ok=True)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    manifest = {"verified_at_utc": datetime.now(timezone.utc).isoformat(), "sources": []}
    for src in SOURCES:
        item = retrieve(src, args.offline)
        manifest["sources"].append(item)
        print(f"Verified {src['id']} {src['file']} {item['sha256']}", flush=True)
    archive = resolve_rottner_archive(args.source_cache, args.offline)
    manifest["remote_only_sources"] = [{**ROTTNER_ARCHIVE, "bytes": archive.stat().st_size,
                                         "storage": "external_cache_not_redistributed", "hash_matches_expected": True}]
    print(f"Verified remote-only {ROTTNER_ARCHIVE['file']} {ROTTNER_ARCHIVE['sha256']}", flush=True)
    (EVIDENCE / "source_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    if args.schema:
        schema()
        print("Wrote evidence/source_schema.json", flush=True)

if __name__ == "__main__":
    main()
