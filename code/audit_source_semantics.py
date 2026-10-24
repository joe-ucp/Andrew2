"""Independent source-semantics audit. Stored zeros are not inferred significance.

This audit reads the pinned workbook's XML and source tables. It distinguishes
published Table S1 labels from illustrative rules applied to stored Table S8
values. It never changes source values or reconstructs unobserved probabilities.
"""
from __future__ import annotations
import argparse
import base64
from collections import Counter, defaultdict
import csv
import hashlib
import json
import io
import math
from pathlib import Path
import re
import struct
from statistics import median
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
from fetch_inputs import ROTTNER_ARCHIVE, resolve_rottner_archive

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
D2 = RAW / "41588_2022_1261_MOESM4_ESM.xlsx"
D2_HASH = "a8f47504edd26d999cb861f146f6fd6aa2ee2ec8368a01e29c808a10238a7690"
NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}

def workbook_transform_audit():
    """Inspect saved import definitions; never execute author M queries or paths."""
    with zipfile.ZipFile(D2) as z:
        workbook = ET.fromstring(z.read("xl/workbook.xml"))
        app = ET.fromstring(z.read("docProps/app.xml"))
        calc = workbook.find("x:calcPr", NS)
        encoded = ET.fromstring(z.read("customXml/item1.xml")).text
        mashup = base64.b64decode(encoded)
        version, package_size = struct.unpack("<II", mashup[:8])
        assert version == 0 and mashup[8:12] == b"PK\x03\x04"
        with zipfile.ZipFile(io.BytesIO(mashup[8:8+package_size])) as package:
            mbytes = package.read("Formulas/Section1.m")
        mtext = mbytes.decode("utf-8-sig")
        queries = {}
        for name in ["sgRNA_hits", '#"RRApaired_screen sgrna_summary (1)"']:
            section = mtext.split("shared " + name + " = let", 1)[1].split(';', 1)[0]
            raw = "shared " + name + " = let" + section + ";"
            # Source paths are not needed to reproduce the column type evidence.
            sanitized = re.sub(r'File.Contents\("[^"]+"\)', 'File.Contents("<source path redacted>")', raw)
            declarations = dict(re.findall(r'\{"([^"]+)", (Int64.Type|type [a-z]+)\}', raw))
            queries[name] = {"column_types": declarations, "query_sha256": hashlib.sha256(raw.encode()).hexdigest(), "query_with_source_path_redacted": sanitized}
        app_values = {e.tag.split('}')[-1]: e.text for e in app if e.tag.split('}')[-1] in {"Application", "AppVersion"}}
        result = {"workbook_sha256": D2_HASH, "custom_xml_part": "customXml/item1.xml", "embedded_package_member": "Formulas/Section1.m", "mashup_version": version, "mashup_package_bytes": package_size, "raw_M_sha256": hashlib.sha256(mbytes).hexdigest(), "producer": app_values, "calcPr_attributes": dict(calc.attrib), "fullPrecision_attribute_present": "fullPrecision" in calc.attrib, "queries": queries,
            "finding": "The saved Table S8 import query explicitly coerces all four p/FDR fields to Int64.Type. The Table S7 query coerces p.low only. This identifies a concrete type-conversion stage consistent with the stored binary cells; it does not establish author intent or recover original values.",
            "documentation": {"integer_conversion": "https://learn.microsoft.com/en-us/powerquery-m/int64-from", "column_conversion": "https://learn.microsoft.com/en-us/powerquery-m/table-transformcolumntypes"}}
    (ROOT / "evidence/workbook_query_transformations.json").write_text(json.dumps(result, indent=2, sort_keys=True), encoding="utf-8")
    return result

def author_archive_audit(cache_dir=None, offline=False):
    path = resolve_rottner_archive(cache_dir, offline)
    with zipfile.ZipFile(path) as z:
        members = [n for n in z.namelist() if not n.endswith('/')]
        rname = next(n for n in members if n.endswith('/CRISPRscreen_EndoCbH1/sgRNA level analysis.R'))
        mname = next(n for n in members if n.endswith('/CRISPRscreen_EndoCbH1/MAGeCK.txt'))
        rbytes = z.read(rname)
        alltext = '\n'.join(z.read(n).decode(errors='replace') for n in members)
    return {"status": "inspected", "source_id": ROTTNER_ARCHIVE["id"], "source_url": ROTTNER_ARCHIVE["url"],
        "storage": "external_cache_not_redistributed", "archive_sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "members": members,
        "selection_script_member": rname, "selection_script_sha256": hashlib.sha256(rbytes).hexdigest(),
        "MAGeCK_command_member": mname,
        "selection_description": "Each control and treatment count >=10; FDR<=0.10; same-coordinate guide matches across replicates; at least two same-direction guides per gene; NDUFS6 removed; separate expression lookup >=1.",
        "source_inputs_not_in_archive": ["RRApaired_screen.sgrna_summary (1).txt", "RRApaired_screen.gene_summary.txt", "LOWHIGH_screen.count.txt", "EndoC_islet_expression lookup table.txt"],
        "rounding_or_integer_cast_function_occurrences": {word: alltext.count(word) for word in ['round(', 'as.integer(', 'as.logical(', 'write.xlsx(', 'write_xlsx(']},
        "scope": "Archived code inspection only; absent processed inputs were not reconstructed and the original pipeline was not rerun."}

def xml_audit():
    assert hashlib.sha256(D2.read_bytes()).hexdigest() == D2_HASH
    with zipfile.ZipFile(D2) as z:
        workbook = ET.fromstring(z.read("xl/workbook.xml"))
        relationships = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships}
        output = {}
        for sheet in workbook.find("x:sheets", NS):
            name = sheet.attrib["name"]
            if name not in {"Table S7-sgRNA_hits", "Table S8-Screen_all sgRNA"}:
                continue
            target = targets[sheet.attrib[f"{{{NS['r']}}}id"]]
            path = target.lstrip("/") if target.startswith("/") else "xl/" + target
            document = ET.fromstring(z.read(path))
            columns = {"K": "p.low", "L": "p.high", "M": "p.twosided", "N": "FDR", "O": "high_in_treatment"}
            records = {label: {"stored_values": Counter(), "cell_types": Counter(), "styles": Counter(), "examples": []} for label in columns.values()}
            for row in document.findall("x:sheetData/x:row", NS):
                if int(row.attrib["r"]) == 1:
                    continue
                for cell in row:
                    col = re.match(r"[A-Z]+", cell.attrib["r"]).group()
                    if col not in columns:
                        continue
                    item = records[columns[col]]
                    value = cell.find("x:v", NS)
                    raw = value.text if value is not None else None
                    item["stored_values"][raw] += 1
                    item["cell_types"][cell.attrib.get("t", "n")] += 1
                    item["styles"][cell.attrib.get("s", "default")] += 1
                    if len(item["examples"]) < 3:
                        item["examples"].append(ET.tostring(cell, encoding="unicode"))
            output[name] = {"xml_part": path, "columns": records}
        return output

def valid_number(x):
    try:
        return math.isfinite(float(x))
    except (TypeError, ValueError):
        return False

def source_tables_audit(compare_revised_outputs=True):
    book = openpyxl.load_workbook(D2, read_only=True, data_only=True)
    hits = {str(r[0]).strip(): str(r[2]).strip() for r in book["Table S1-Screen hits "].iter_rows(min_row=2, max_col=3, values_only=True) if r[0]}
    expression_rows = defaultdict(list)
    for rowno, r in enumerate(book["Table S5 - RNASeq"].iter_rows(min_row=2, max_col=3, values_only=True), 2):
        if r[1]:
            expression_rows[str(r[1]).strip()].append({"GeneID": r[0], "baseMean": r[2], "row": rowno})
    expression = {g: float(rows[0]["baseMean"]) for g, rows in expression_rows.items() if len(rows) == 1 and valid_number(rows[0]["baseMean"])}
    guides = defaultdict(lambda: defaultdict(dict))
    s8_match = defaultdict(list)
    for rowno, r in enumerate(book["Table S8-Screen_all sgRNA"].iter_rows(min_row=2, values_only=True), 2):
        match = re.fullmatch(r"(.*)_r([01])", str(r[0]).strip())
        if not match or not r[1] or not all(valid_number(r[i]) for i in (2, 3, 6, 13)):
            continue
        key, rep = match.groups()
        g = str(r[1]).strip()
        assert key not in guides[g][rep], (g, rep, key)
        item = {"gene": g, "replicate": "r"+rep, "guide_key": key, "row": rowno, "LFC": float(r[6]), "stored_FDR": float(r[13]), "control_count": float(r[2]), "treatment_count": float(r[3]), "p_low": float(r[10]), "p_high": float(r[11]), "p_twosided": float(r[12])}
        guides[g][rep][key] = item
        signature = (key.split("_")[0], g, float(r[2]), float(r[3]), float(r[6]))
        s8_match[signature].append(item)
    linkage = []
    ambiguous_linkage = []
    for rowno, r in enumerate(book["Table S7-sgRNA_hits"].iter_rows(min_row=2, values_only=True), 2):
        signature = (str(r[0]).strip(), str(r[1]).strip(), float(r[2]), float(r[3]), float(r[6]))
        found = s8_match.get(signature, [])
        if len(found) == 1:
            s8 = found[0]
            linkage.append({"gene": str(r[1]).strip(), "coordinate": str(r[0]).strip(), "guide_key": s8["guide_key"], "S7_row": rowno, "S8_row": s8["row"], "replicate": s8["replicate"], "LFC": s8["LFC"], "control_count": s8["control_count"], "treatment_count": s8["treatment_count"], "S7_FDR": float(r[13]), "S8_stored_FDR": s8["stored_FDR"], "S7_p_high": float(r[11]), "S8_stored_p_high": s8["p_high"], "S7_p_twosided": float(r[12]), "S8_stored_p_twosided": s8["p_twosided"]})
        else:
            ambiguous_linkage.append({"S7_row": rowno, "gene": str(r[1]).strip(), "number_of_S8_matches": len(found)})
    book.close()
    measured_genes = set(guides)
    with (RAW / "Supplementary_Data_09.CRISPR-KO-screen_results.csv").open(encoding="utf-8-sig", newline="") as f:
        survival_rows = list(csv.DictReader(f))
    names = Counter(str(r["Gene"]).strip() for r in survival_rows)
    survival = {str(r["Gene"]).strip(): r for r in survival_rows if names[str(r["Gene"]).strip()] == 1 and not str(r["Gene"]).strip().startswith("hsa-") and all(valid_number(r[c]) for c in ("sgRNA number", "Cytokine survival beta", "Cytokine survival fdr")) and float(r["sgRNA number"]) >= 3}
    records = []
    for g in sorted(set(guides) | set(survival)):
        reps = guides[g]
        common_keys = set(reps["0"]) & set(reps["1"])
        covered = {key for key in common_keys if all(reps[rep][key]["control_count"] + reps[rep][key]["treatment_count"] >= 20 for rep in ("0", "1"))}
        pos = {rep: {key for key, r in reps[rep].items() if r["stored_FDR"] == 0 and r["LFC"] > 0} for rep in ("0", "1")}
        neg = {rep: {key for key, r in reps[rep].items() if r["stored_FDR"] == 0 and r["LFC"] < 0} for rep in ("0", "1")}
        positive_paired = pos["0"] & pos["1"]
        negative_paired = neg["0"] & neg["1"]
        both_counts_ten = {key for key in common_keys if all(reps[rep][key]["control_count"] >= 10 and reps[rep][key]["treatment_count"] >= 10 for rep in ("0", "1"))}
        u0 = g in survival and g not in {"EGFP", "LacZ", "luciferase"} and len(covered) >= 3
        raw_expression_eligible = g in expression and expression[g] >= 10
        published = hits.get(g, "Not a published hit")
        stored_rule = all(len(pos[rep]) >= 2 for rep in ("0", "1"))
        stored_negative_rule = all(len(neg[rep]) >= 2 for rep in ("0", "1"))
        legacy_direction_pass = published == "Not a published hit" or (stored_rule if published == "Decrease" else stored_negative_rule)
        in_legacy_U = u0 and raw_expression_eligible and legacy_direction_pass
        a = g in survival and float(survival[g]["Cytokine survival beta"]) > 0 and float(survival[g]["Cytokine survival fdr"]) < .1
        screen_status = {"Decrease": "Published lower-content evidence", "Increase": "Published higher-content evidence"}.get(published, "No published directional call" if g in measured_genes else "Screen measurement unavailable")
        records.append({"gene": g, "published_content_label": published, "screen_evidence_status": screen_status, "baseMean_exact_unique": expression.get(g), "expression_rows": len(expression_rows.get(g, [])), "covered_insulin_guides": len(covered), "in_U0": u0, "in_revised_primary_U": u0 and raw_expression_eligible, "in_legacy_primary_U": in_legacy_U, "survival_favoring": a, "stored_zero_positive_r0": len(pos["0"]), "stored_zero_positive_r1": len(pos["1"]), "stored_zero_negative_r0": len(neg["0"]), "stored_zero_negative_r1": len(neg["1"]), "illustrative_positive_rule": stored_rule, "illustrative_negative_rule": stored_negative_rule, "paired_identical_positive_zero_guides": len(positive_paired), "paired_identical_negative_zero_guides": len(negative_paired), "paired_positive_zero_guides_with_each_bin_count_ge10": len(positive_paired & both_counts_ten), "paired_negative_zero_guides_with_each_bin_count_ge10": len(negative_paired & both_counts_ten), "positive_paired_keys": ";".join(sorted(positive_paired)), "negative_paired_keys": ";".join(sorted(negative_paired)), "r0_median_LFC": median(r["LFC"] for r in reps["0"].values()) if reps["0"] else None, "r1_median_LFC": median(r["LFC"] for r in reps["1"].values()) if reps["1"] else None})
    primary = [r for r in records if r["in_legacy_primary_U"]]
    candidates14 = [r for r in primary if r["survival_favoring"] and r["published_content_label"] == "Decrease"]
    nonpublished_positive = [r for r in primary if r["published_content_label"] == "Not a published hit" and r["illustrative_positive_rule"]]
    def prevalence(rows):
        return {"n": len(rows), "any_zero_negative_guide_either_replicate": sum(r["stored_zero_negative_r0"] + r["stored_zero_negative_r1"] > 0 for r in rows), "at_least_one_zero_negative_guide_each_replicate": sum(r["stored_zero_negative_r0"] >= 1 and r["stored_zero_negative_r1"] >= 1 for r in rows), "at_least_two_zero_negative_guides_each_replicate": sum(r["illustrative_negative_rule"] for r in rows), "at_least_two_identical_positive_zero_guides_across_replicates": sum(r["paired_identical_positive_zero_guides"] >= 2 for r in rows)}
    diagnostics = {"legacy_primary_n": len(primary), "nonpublished_primary_positive_stored_zero_rule_count": len(nonpublished_positive), "nonpublished_primary_positive_stored_zero_rule_genes": [r["gene"] for r in nonpublished_positive], "published_primary_decrease_count": sum(r["published_content_label"] == "Decrease" for r in primary), "opposing_stored_zero_guide_prevalence": {"all_primary_genes": prevalence(primary), "published_primary_decreases": prevalence([r for r in primary if r["published_content_label"] == "Decrease"]), "legacy14_candidates": prevalence(candidates14)}, "legacy14_candidates": candidates14, "TOMM6_and_NUPL1": [r for r in records if r["gene"] in {"TOMM6", "NUPL1"}], "special_guide_rows": {g: [r for rep in guides[g].values() for r in rep.values()] for g in {r["gene"] for r in candidates14} | {"TOMM6", "NUPL1"}}, "S7_S8_linkage": {"unambiguous_rows": len(linkage), "ambiguous_or_missing": ambiguous_linkage, "FDR_S8_equals_nearest_integer_S7_count": sum(r["S8_stored_FDR"] == round(r["S7_FDR"]) for r in linkage), "S7_FDR_greater_than_0_1_with_S8_zero": [r for r in linkage if r["S7_FDR"] > .1 and r["S8_stored_FDR"] == 0], "interpretation": "Matching uses exact source gene, genomic coordinate, control count, treatment count and LFC. Numeric rounding compatibility is an observation, not proof of the export mechanism. S7 is selected-guide evidence and cannot repair the complete S8 denominator."}}
    for name, rows in (("semantic_gene_diagnostics.csv", records), ("S7_S8_semantic_linkage.csv", linkage)):
        with (ROOT / "evidence" / name).open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)
    diagnostics["independent_revised_catalogue_checks"] = revised_catalogue_checks(records, compare_revised_outputs)
    candidate_s7 = []
    for gene in sorted({r["gene"] for r in candidates14} | {"NUPL1", "TOMM6"}):
        rows = [r for r in linkage if r["gene"] == gene]
        item = {"gene": gene, "evidence_scope": "Selected Table S7 subset; not complete genome-wide guide probabilities", "all_source_rows_unambiguously_linked": True}
        for rep in ("r0", "r1"):
            subset = [r for r in rows if r["replicate"] == rep]
            item[f"S7_rows_{rep}"] = len(subset)
            item[f"S8_rows_{rep}"] = len(guides[gene][rep[-1]])
            item[f"S7_covers_all_S8_rows_{rep}"] = len(subset) == len(guides[gene][rep[-1]])
            item[f"source_s7_positive_fdr_lt010_{rep}"] = sum(r["LFC"] > 0 and r["S7_FDR"] < .1 for r in subset)
            item[f"source_s7_positive_fdr_le010_{rep}"] = sum(r["LFC"] > 0 and r["S7_FDR"] <= .1 for r in subset)
            item[f"source_s7_negative_fdr_lt010_{rep}"] = sum(r["LFC"] < 0 and r["S7_FDR"] < .1 for r in subset)
        poskeys = [{r["guide_key"] for r in rows if r["replicate"] == rep and r["LFC"] > 0 and r["S7_FDR"] < .1} for rep in ("r0", "r1")]
        item["source_s7_positive_fdr_lt010_sameguide"] = len(poskeys[0] & poskeys[1])
        item["source_s7_positive_fdr_lt010_sameguide_keys"] = ";".join(sorted(poskeys[0] & poskeys[1]))
        item["source_s7_any_FDR_exactly010"] = any(r["S7_FDR"] == .1 for r in rows)
        item["source_s7_three_positive_each_replicate"] = all(item[f"source_s7_positive_fdr_lt010_{rep}"] >= 3 for rep in ("r0", "r1"))
        item["source_s7_three_identical_positive_guides"] = item["source_s7_positive_fdr_lt010_sameguide"] >= 3
        candidate_s7.append(item)
    with (ROOT / "evidence/candidate_S7_selected_guide_evidence.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(candidate_s7[0])); writer.writeheader(); writer.writerows(candidate_s7)
    diagnostics["candidate_S7_selected_subset_evidence"] = {"candidates": candidate_s7, "all_S7_exact_FDR_010_ties": sum(r["S7_FDR"] == .1 for r in linkage), "scope": "Full-precision source values for the published selected-guide subset only. Matches establish source alignment; they are not an independent experiment, reconstructed genome-wide significance, or complete evidence of nonsignificant/opposing guides."}
    return diagnostics


def revised_catalogue_checks(records, compare_outputs=True):
    """Compare raw-source reconstruction to V2 outputs without importing analysis."""
    def counts(rows):
        return {"n": len(rows), "A": sum(r["survival_favoring"] for r in rows),
                "L": sum(r["published_content_label"] == "Decrease" for r in rows),
                "overlap": sum(r["survival_favoring"] and r["published_content_label"] == "Decrease" for r in rows)}
    result = {"comparison_status": "raw_reconstruction_completed" if compare_outputs else "output_comparisons_explicitly_deferred", "analyses": {}}
    for name, flag in (("primary", "in_revised_primary_U"), ("no_expression", "in_U0")):
        selected = [r for r in records if r[flag]]
        audit = counts(selected)
        audit["lower_content_candidates"] = sorted(r["gene"] for r in selected if r["survival_favoring"] and r["published_content_label"] == "Decrease")
        path = ROOT / "results" / (name + "_genes.csv")
        if compare_outputs and path.exists():
            with path.open(newline="", encoding="utf-8") as f:
                actual = {r["gene"]: r for r in csv.DictReader(f)}
            expected = {r["gene"]: r for r in selected}
            audit["output_membership_identical"] = set(actual) == set(expected)
            audit["output_A_L_identical"] = set(actual) == set(expected) and all((actual[g]["A"] == "True") == r["survival_favoring"] and (actual[g]["L"] == "True") == (r["published_content_label"] == "Decrease") for g, r in expected.items())
            audit["output_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
            assert audit["output_membership_identical"] and audit["output_A_L_identical"], (name, audit)
        result["analyses"][name] = audit
    candidates = [r for r in records if r["survival_favoring"]]
    result["candidate_axes"] = {"total": len(candidates), "common_coverage": sum(r["in_U0"] for r in candidates),
        "primary_eligible": sum(r["in_revised_primary_U"] for r in candidates),
        "screen_evidence_status_counts": dict(Counter(r["screen_evidence_status"] for r in candidates)),
        "measured_but_coverage_excluded_no_call": sum(not r["in_U0"] and r["screen_evidence_status"] == "No published directional call" for r in candidates)}
    path = ROOT / "results/candidate_table_427.csv"
    if compare_outputs and path.exists():
        with path.open(newline="", encoding="utf-8") as f:
            actual = {r["gene"]: r for r in csv.DictReader(f)}
        expected = {r["gene"]: r for r in candidates}
        if actual and "screen_evidence_status" in next(iter(actual.values())):
            same = set(actual) == set(expected) and all(actual[g]["screen_evidence_status"] == r["screen_evidence_status"] and (actual[g]["in_U0"] == "True") == r["in_U0"] and (actual[g]["primary_eligible"] == "True") == r["in_revised_primary_U"] for g, r in expected.items())
            result["candidate_axes"]["output_membership_and_axes_identical"] = same
            result["candidate_axes"]["output_sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
            assert same, "V2 candidate membership or axes disagree with raw reconstruction"
    return result

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xml-only", action="store_true")
    ap.add_argument("--offline", action="store_true", help="Require the verified external author-code cache; never download")
    ap.add_argument("--source-cache", help="External author-code cache directory; same option as code/fetch_inputs.py")
    ap.add_argument("--defer-revised-checks", action="store_true", help="Initial pipeline stage only: compute raw diagnostics/S7 but explicitly defer comparisons to not-yet-built V2 output tables")
    args = ap.parse_args()
    manifest = json.loads((ROOT / "evidence/input_manifest.json").read_text())
    for entry in manifest["inputs"]:
        if entry["id"] in {"D1", "D2"}:
            assert hashlib.sha256((ROOT / entry["path"]).read_bytes()).hexdigest() == entry["sha256"], entry["id"]
    result = {"source_file": "data/raw/41588_2022_1261_MOESM4_ESM.xlsx", "sha256": D2_HASH, "xml": xml_audit(), "semantic_status": "QUALIFIED_CATALOGUE_ANALYSIS_GUIDE_SIGNIFICANCE_UNRESOLVED", "review_disposition": "Post-review V2 retains published Table S1 decisions and independently reports stored-value diagnostics. The binary guide fields cannot establish the original guide-significance threshold.", "scientific_validation_superseded": "V1 independent agreement validates arithmetic and implementation only. It does not establish calibrated p/FDR meaning, hit-selection equivalence, biological replication, or the original scientific interpretation.", "interpretive_rule": "A stored 0 is an observed encoded value; this audit does not interpret it as a validated FDR or p-value."}
    if not args.xml_only:
        result["diagnostics"] = source_tables_audit(compare_revised_outputs=not args.defer_revised_checks)
    result["saved_import_pipeline"] = workbook_transform_audit()
    result["archived_author_code"] = author_archive_audit(args.source_cache, args.offline)
    result["claim_permissions"] = {
        "V1_S8_FDR_below_010_independent_guide_significance": "REJECTED",
        "complete_genomewide_continuous_q_values_reconstructed": False,
        "published_Table_S1_catalogue_comparison": "PERMITTED_WITH_QUALIFICATION",
        "selected_Table_S7_continuous_values": "SOURCE_SUPPORTED_SELECTED_SUBSET_ONLY",
        "scope": "Numerical agreement does not validate the rejected guide-significance claim. Published catalogue labels and selected source probabilities retain their explicitly stated limits; source values are unchanged."
    }
    result["source_hashes"] = {str(p.relative_to(ROOT)).replace("\\", "/"): hashlib.sha256(p.read_bytes()).hexdigest() for p in [D2, RAW / "Supplementary_Data_09.CRISPR-KO-screen_results.csv"] + list((ROOT / "data/source_documents").glob("*")) if p.is_file() and p.name != ROTTNER_ARCHIVE["file"]}
    result["remote_source_hashes"] = {ROTTNER_ARCHIVE["id"]: result["archived_author_code"]["archive_sha256"]}
    result["audit_script_sha256"] = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    result["derived_evidence_hashes"] = {str(p.relative_to(ROOT)).replace("\\", "/"): hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT / "evidence/semantic_gene_diagnostics.csv", ROOT / "evidence/S7_S8_semantic_linkage.csv", ROOT / "evidence/candidate_S7_selected_guide_evidence.csv"] if p.exists()}
    (ROOT / "evidence").mkdir(exist_ok=True)
    (ROOT / "evidence/source_semantics_audit.json").write_text(json.dumps(result, indent=2, sort_keys=True), encoding="utf-8")
    print(json.dumps({name: {col: {"distinct": len(val["stored_values"]), "zero": val["stored_values"].get("0", 0), "one": val["stored_values"].get("1", 0), "types": val["cell_types"]} for col, val in content["columns"].items()} for name, content in result["xml"].items()}, indent=2))
    if "diagnostics" in result:
        print(json.dumps({k: v for k, v in result["diagnostics"].items() if k in {"legacy_primary_n", "nonpublished_primary_positive_stored_zero_rule_count", "opposing_stored_zero_guide_prevalence", "S7_S8_linkage"}}, indent=2))

if __name__ == "__main__":
    main()
